"""Non-destructive inventory and located extraction. Python 3.10+."""
import argparse
import csv
import hashlib
import importlib.metadata
import json
import os
import sys
from datetime import datetime, timezone
from pathlib import Path


def digest(path):
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def extract(path):
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".json", ".csv", ".tsv"}:
        text = path.read_text(encoding="utf-8-sig", errors="strict")
        if suffix in {".csv", ".tsv"}:
            rows = csv.reader(text.splitlines(keepends=True), delimiter="\t" if suffix == ".tsv" else ",")
            units = [{"location": {"row": i}, "text": json.dumps(row, ensure_ascii=False)}
                     for i, row in enumerate(rows, 1)]
            return units, "stdlib.csv", []
        if suffix == ".json":
            json.loads(text)  # Fail visibly on malformed JSON.
        return [{"location": {"line": i}, "text": line}
                for i, line in enumerate(text.splitlines(), 1)], "stdlib.utf8", []
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError("encrypted PDF: obtain an authorized decrypted copy")
        units = [{"location": {"page": i}, "text": page.extract_text() or ""}
                 for i, page in enumerate(reader.pages, 1)]
        issues = [f"page {u['location']['page']}: empty text; inspect/OCR" for u in units if not u["text"].strip()]
        return units, "pypdf@" + importlib.metadata.version("pypdf"), issues
    if suffix == ".docx":
        from docx import Document
        doc = Document(path)
        units = [{"location": {"paragraph": i}, "text": p.text} for i, p in enumerate(doc.paragraphs, 1)]
        for ti, table in enumerate(doc.tables, 1):
            for ri, row in enumerate(table.rows, 1):
                units.append({"location": {"table": ti, "row": ri}, "text": json.dumps([c.text for c in row.cells], ensure_ascii=False)})
        return units, "python-docx@" + importlib.metadata.version("python-docx"), ["headers, textboxes, nested tables and embedded objects require visual review"]
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False)
        units = []
        try:
            for sheet in book:
                for row in sheet:
                    for cell in row:
                        if cell.value is not None:
                            units.append({"location": {"sheet": sheet.title, "cell": cell.coordinate}, "text": str(cell.value)})
        finally:
            book.close()
        return units, "openpyxl@" + importlib.metadata.version("openpyxl"), ["formulas preserved, not recalculated; inspect charts, merges and embedded objects"]
    return [], "none", ["needs format-specific extraction/OCR/transcription"]


def prepare(source, output, expected=None):
    source, output = Path(source).resolve(), Path(output).resolve()
    if not source.is_dir():
        raise ValueError("input must be a directory")
    if output == source or source in output.parents or output in source.parents:
        raise ValueError("input and output must be separate, non-nested directories")
    if output.exists() and any(output.iterdir()):
        raise ValueError("output must be new or empty; do not overwrite evidence")
    expected = [] if expected is None else expected
    if not isinstance(expected, list) or any(not isinstance(x, str) for x in expected):
        raise ValueError("expected must be a JSON array of relative paths")
    for item in expected:
        p = Path(item)
        if p.is_absolute() or ".." in p.parts:
            raise ValueError("expected paths must stay inside input")
    output.mkdir(parents=True, exist_ok=True)
    manifest = {"schema_version": 1, "source_root": str(source), "generated_at": datetime.now(timezone.utc).isoformat(),
                "filesystem_time_zone": "UTC", "python": sys.version.split()[0], "files": [], "missing": [], "issues": []}
    evidence, hashes, seen = [], {}, set()
    paths = []
    for directory, dirs, names in os.walk(source, followlinks=False):
        for name in list(dirs):
            candidate = Path(directory) / name
            if candidate.is_symlink() or getattr(candidate, "is_junction", lambda: False)():
                manifest["issues"].append({"path": candidate.relative_to(source).as_posix(), "reason": "directory link skipped"})
                dirs.remove(name)
        paths.extend(Path(directory) / name for name in names)
    for path in sorted(paths):
        rel = path.relative_to(source).as_posix()
        if path.is_symlink() or getattr(path, "is_junction", lambda: False)():
            manifest["issues"].append({"path": rel, "reason": "link skipped"})
            continue
        if not path.is_file():
            continue
        item = {"path": rel, "status": "error", "issues": []}
        manifest["files"].append(item)
        seen.add(rel)
        try:
            item.update(sha256=digest(path), bytes=path.stat().st_size,
                        mtime_utc=datetime.fromtimestamp(path.stat().st_mtime, timezone.utc).isoformat())
            if item["sha256"] in hashes:
                item["duplicate_of"] = hashes[item["sha256"]]
            else:
                hashes[item["sha256"]] = rel
            units, tool, issues = extract(path)
            item.update(tool=tool, issues=issues, status="needs_extraction" if tool == "none" else "extracted")
            if issues and tool != "none":
                item["status"] = "review_required"
            for n, unit in enumerate(units, 1):
                evidence.append({"id": f"{rel}#{n}", "source": rel, "source_sha256": item["sha256"], **unit})
            item["units"] = len(units)
            if not units and tool != "none":
                item["status"] = "review_required"
                item["issues"].append("empty extraction")
        except Exception as exc:
            item["issues"].append(f"{type(exc).__name__}: {exc}")
    manifest["missing"] = sorted(set(x.replace("\\", "/") for x in expected) - seen)
    for name, data in [("manifest.json", manifest), ("evidence.json", {"schema_version": 1, "units": evidence})]:
        (output / name).write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return manifest


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--expected", type=Path)
    args = parser.parse_args()
    try:
        expected = json.loads(args.expected.read_text(encoding="utf-8-sig")) if args.expected else None
        result = prepare(args.input, args.output, expected)
    except (ValueError, OSError) as exc:
        parser.exit(2, f"{exc}\n")
    print(json.dumps({"files": len(result["files"]), "missing": result["missing"], "statuses": [x["status"] for x in result["files"]]}))
    return 1 if result["missing"] or result["issues"] or any(x["status"] == "error" for x in result["files"]) else 0


if __name__ == "__main__":
    raise SystemExit(main())
